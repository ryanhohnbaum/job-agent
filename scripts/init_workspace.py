"""Create a JobAgent workspace: folders, tracker workbook, and starter config."""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from schema import APPLICATION_HEADERS, CORRESPONDENCE_HEADERS, SEARCH_ACTIVITY_HEADERS, STATUSES

PLUGIN_ROOT = Path(__file__).resolve().parent.parent


def style_sheet(ws, header_row: list[str]) -> None:
    ws.append(header_row)
    fill = PatternFill("solid", fgColor="1F3864")
    for idx, cell in enumerate(ws[1], start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        ws.column_dimensions[get_column_letter(idx)].width = max(14, min(40, len(header_row[idx - 1]) + 4))
    ws.freeze_panes = "A2"


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="Create a JobAgent workspace.")
    parser.add_argument("--root", required=True, help="Folder to create/use as the workspace.")
    args = parser.parse_args(argv)

    root = Path(args.root).expanduser().resolve()
    for sub in ("Applications", "Templates", "Config", "Backups", "Logs"):
        (root / sub).mkdir(parents=True, exist_ok=True)

    book = root / "Job_Search_Log.xlsx"
    if book.exists():
        print(f"KEPT existing tracker: {book}")
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Applications"
        style_sheet(ws, APPLICATION_HEADERS)
        style_sheet(wb.create_sheet("Correspondence"), CORRESPONDENCE_HEADERS)
        style_sheet(wb.create_sheet("Search Activity"), SEARCH_ACTIVITY_HEADERS)
        guide = wb.create_sheet("Status Guide")
        guide.append(["Status"])
        for status in STATUSES:
            guide.append([status])
        wb.save(book)
        print(f"CREATED tracker: {book}")

    for name in ("profile.yaml", "search_rules.yaml"):
        dest = root / "Config" / name
        if dest.exists():
            print(f"KEPT existing config: {dest}")
        else:
            shutil.copy2(PLUGIN_ROOT / "templates" / f"{name}.template", dest)
            print(f"CREATED config: {dest}")

    ignore = root / ".gitignore"
    if not ignore.exists():
        ignore.write_text("# Personal job-search data. Never commit this workspace.\n*\n", encoding="utf-8")

    print(f"WORKSPACE READY: {root}")


if __name__ == "__main__":
    main()
