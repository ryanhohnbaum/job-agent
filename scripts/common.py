"""Shared helpers for JobAgent scripts. Cross-platform (Windows, macOS, Linux)."""
from __future__ import annotations

import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

MARKER = Path("Config") / "profile.yaml"

try:  # keep Unicode output from crashing on legacy Windows consoles
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def jobagent_root() -> Path:
    """Workspace root: $JOB_AGENT_ROOT, else the nearest parent of the cwd that holds
    Config/profile.yaml, else ~/JobAgent."""
    configured = os.environ.get("JOB_AGENT_ROOT", "").strip()
    if configured:
        return Path(configured).expanduser()
    here = Path.cwd().resolve()
    for candidate in (here, *here.parents):
        if (candidate / MARKER).exists():
            return candidate
    return Path.home() / "JobAgent"


def config_dir() -> Path:
    return jobagent_root() / "Config"


def workbook_path() -> Path:
    return jobagent_root() / "Job_Search_Log.xlsx"


def workbook_lock_path() -> Path:
    return jobagent_root() / "~$Job_Search_Log.xlsx"


def load_yaml(path: Path) -> dict:
    import yaml

    if not path.exists():
        return {}
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {}


def load_profile() -> dict:
    return load_yaml(config_dir() / "profile.yaml")


def require_workbook_writable() -> Path:
    book = workbook_path()
    if not book.exists():
        raise FileNotFoundError(f"{book} not found. Run the setup skill (or `jobagent.py init`) first.")
    if workbook_lock_path().exists():
        raise RuntimeError("Job_Search_Log.xlsx appears open in Excel. Close it and rerun.")
    return book


def headers(ws, required: Iterable[str] = ()) -> dict[str, int]:
    result = {
        str(cell.value).strip(): index
        for index, cell in enumerate(ws[1], start=1)
        if cell.value not in (None, "")
    }
    missing = [name for name in required if name not in result]
    if missing:
        raise RuntimeError(f"Worksheet {ws.title!r} is missing column(s): {', '.join(missing)}")
    return result


def first_blank_row(ws, key_column: int, start_row: int = 2) -> int:
    row = start_row
    while ws.cell(row, key_column).value not in (None, ""):
        row += 1
    return row


def iter_value_rows(ws, start_row: int = 2):
    for row_number, values in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start=start_row):
        yield row_number, values


def row_value(values: tuple, column: int):
    return values[column - 1] if column <= len(values) else None


def normalized(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def clean_name(value: str, limit: int = 80) -> str:
    value = re.sub(r"[^\w]+", "_", value.strip(), flags=re.UNICODE).strip("_")
    return value[:limit]


def create_workbook_backup(action: str) -> Path:
    book = workbook_path()
    backup_dir = jobagent_root() / "Backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup = backup_dir / f"Job_Search_Log_before_{action}_{stamp}.xlsx"
    shutil.copy2(book, backup)
    return backup


def save_workbook_verified(wb, action: str, verifier: Callable[[Path], None]) -> tuple[Path, Path]:
    """Backup, save to a temp file, verify the temp file, then atomically replace the workbook."""
    book = workbook_path()
    backup = create_workbook_backup(action)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    temp = book.with_name(f"{book.stem}.{action}.{stamp}.tmp{book.suffix}")
    try:
        wb.save(temp)
        wb.close()
        verifier(temp)
        temp.replace(book)
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        temp.unlink(missing_ok=True)
        raise
    return book, backup


def verify_application_values(path: Path, application_id: str, expected: dict[str, object]) -> None:
    from openpyxl import load_workbook

    check = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = check["Applications"]
        h = headers(ws, ["Application ID", *expected.keys()])
        rows = [
            r for r, v in iter_value_rows(ws)
            if str(row_value(v, h["Application ID"]) or "").strip() == application_id
        ]
        if len(rows) != 1:
            raise RuntimeError(f"Verification expected one row for {application_id}; found {len(rows)}.")
        row = rows[0]
        for name, value in expected.items():
            actual = ws.cell(row, h[name]).value
            if isinstance(value, datetime):
                if not isinstance(actual, datetime) or actual.date() != value.date():
                    raise RuntimeError(f"Verification failed for {application_id} column {name}.")
            elif str(actual or "") != str(value or ""):
                raise RuntimeError(
                    f"Verification failed for {application_id} column {name}: "
                    f"expected {value!r}, found {actual!r}."
                )
    finally:
        check.close()
