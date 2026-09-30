from __future__ import annotations

import argparse
import base64
import html as html_lib
import json
import os
import re
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any

from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

try:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from common import jobagent_root, load_profile  # noqa: E402

ROOT = jobagent_root()
CONFIG = ROOT / "Config"
LOG_FILE = ROOT / "Job_Search_Log.xlsx"
TOKEN_FILE = CONFIG / "gmail_token.json"
THREAD_MAP_FILE = CONFIG / "gmail_threads.json"
PROCESSED_FILE = CONFIG / "gmail_processed.json"
PENDING_FILE = CONFIG / "gmail_pending.json"
SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]

ARCHIVABLE_KINDS = {
    "Application Confirmation",
    "Recruiter Outreach",
    "Interview Invite",
    "Screening Request",
    "Scheduling",
    "Assessment",
    "Follow-up",
    "Rejection",
    "Offer",
    "Withdrawal",
}

TERMINAL_KINDS = {"Rejection", "Offer", "Withdrawal"}


def norm(text: Any) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(text or "").lower())).strip()


def compact_norm(text: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(text or "").lower())


def safe_filename(text: Any) -> str:
    text = re.sub(r'[<>:"/\\|?*]+', "", str(text or ""))
    text = re.sub(r"\s+", "_", text.strip())
    return text[:90] or "Email"


def decode_b64(data: str | None) -> str:
    if not data:
        return ""
    return base64.urlsafe_b64decode(data.encode("ascii")).decode("utf-8", errors="replace")


def walk_parts(part: dict):
    yield part
    for child in part.get("parts", []) or []:
        yield from walk_parts(child)


def strip_tags(s: str) -> str:
    s = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", s or "")
    s = re.sub(r"(?is)<br\s*/?>", "\n", s)
    s = re.sub(r"(?is)</p\s*>", "\n", s)
    s = re.sub(r"(?s)<[^>]+>", " ", s)
    return re.sub(r"\s+", " ", html_lib.unescape(s)).strip()


def message_content(payload: dict) -> tuple[str, str, str]:
    html_body = ""
    text_body = ""
    for part in walk_parts(payload):
        mime = part.get("mimeType", "")
        data = (part.get("body") or {}).get("data")
        if not data:
            continue
        decoded = decode_b64(data)
        if mime == "text/html" and not html_body:
            html_body = decoded
        elif mime == "text/plain" and not text_body:
            text_body = decoded
    visible = text_body.strip() or strip_tags(html_body)
    if not html_body:
        escaped = html_lib.escape(text_body or "")
        html_body = f"<pre style='white-space:pre-wrap;font-family:Segoe UI,Arial,sans-serif'>{escaped}</pre>"
    return html_body, text_body, visible


def get_headers(payload: dict) -> dict[str, str]:
    return {h.get("name", "").lower(): h.get("value", "") for h in payload.get("headers", [])}


def self_addresses() -> set[str]:
    """The user's own addresses, from Config/profile.yaml (candidate.email + gmail.extra_addresses)."""
    profile = load_profile()
    addrs = {str((profile.get("candidate") or {}).get("email") or "").strip().lower()}
    addrs |= {str(a).strip().lower() for a in ((profile.get("gmail") or {}).get("extra_addresses") or [])}
    return {a for a in addrs if a}


def direction_from_headers(headers: dict[str, str]) -> str:
    sender = headers.get("from", "").lower()
    if any(addr in sender for addr in self_addresses()):
        return "Outgoing"
    return "Incoming"


def email_datetime(headers: dict[str, str]) -> datetime:
    raw = headers.get("date")
    try:
        dt = parsedate_to_datetime(raw)
        return dt.astimezone() if dt.tzinfo else dt
    except Exception:
        return datetime.now().astimezone()


def get_service():
    if not TOKEN_FILE.exists():
        raise FileNotFoundError(f"Missing {TOKEN_FILE}. Run Gmail auth first.")
    creds = Credentials.from_authorized_user_file(str(TOKEN_FILE), SCOPES)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
        atomic_write_text(TOKEN_FILE, creds.to_json())
    return build("gmail", "v1", credentials=creds)


def is_gmail_retryable_error(exc: Exception) -> bool:
    """Return True for Gmail/API/transient transport errors worth retrying."""
    text = str(exc).lower()
    if isinstance(exc, HttpError):
        status = getattr(exc.resp, "status", None)
        return status in (403, 429, 500, 502, 503, 504) and (
            "ratelimitexceeded" in text
            or "rate limit" in text
            or "quota exceeded" in text
            or "user-rate-limit-exceeded" in text
            or "backend error" in text
            or "internal error" in text
            or "service unavailable" in text
        )
    retryable_markers = (
        "winerror 10054",
        "forcibly closed by the remote host",
        "connection reset",
        "connection aborted",
        "remote end closed connection",
        "temporarily unavailable",
        "timed out",
        "timeout",
        "ssl",
        "tls",
        "broken pipe",
    )
    return isinstance(exc, (OSError, ConnectionError, TimeoutError)) or any(m in text for m in retryable_markers)


def execute_gmail_request(request, *, label: str, max_retries: int = 5, base_sleep: float = 8.0):
    """Execute a Gmail API request with backoff for quota and transient network errors."""
    attempt = 0
    while True:
        try:
            return request.execute()
        except Exception as exc:
            if not is_gmail_retryable_error(exc) or attempt >= max_retries:
                raise
            wait = min(120.0, base_sleep * (2 ** attempt))
            print(f"RETRYABLE GMAIL ERROR {label}: {exc}")
            print(f"Sleeping {wait:.1f}s then retrying ({attempt + 1}/{max_retries})")
            time.sleep(wait)
            attempt += 1


def load_json(path: Path, default):
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return default
    return default


def atomic_write_text(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + f".{datetime.now():%Y%m%d%H%M%S%f}.tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(path)


def save_json(path: Path, data):
    atomic_write_text(path, json.dumps(data, indent=2, sort_keys=True))


def ws_headers(ws) -> dict[str, int]:
    return {str(c.value).strip(): i for i, c in enumerate(ws[1], 1) if c.value not in (None, "")}


def find_col(h: dict[str, int], *names: str) -> int | None:
    direct = {norm(k): v for k, v in h.items()}
    for name in names:
        v = direct.get(norm(name))
        if v:
            return v
    return None


def require_col(h: dict[str, int], *names: str) -> int:
    v = find_col(h, *names)
    if not v:
        raise KeyError(f"Missing required column; tried: {', '.join(names)}")
    return v


def ensure_col(ws, h: dict[str, int], name: str) -> int:
    col = find_col(h, name)
    if col:
        return col
    col = ws.max_column + 1
    ws.cell(1, col).value = name
    h[name] = col
    return col


def first_blank_row(ws, key_col: int = 1) -> int:
    r = 2
    while ws.cell(r, key_col).value not in (None, ""):
        r += 1
    return r


@dataclass
class Classification:
    kind: str
    confidence: float
    reason: str


def contains_any(text_n: str, phrases: tuple[str, ...] | list[str]) -> str | None:
    for p in phrases:
        pn = norm(p)
        if pn and pn in text_n:
            return p
    return None


def classify(subject: str, body: str, sender: str = "") -> Classification:
    """Content-aware employment email classifier.

    Important precedence rule: outcome/disposition beats courtesy language.
    Many rejection emails begin with "thank you for applying" or mention the
    hiring team; those must still classify as Rejection when the body says the
    candidate will not move forward.
    """
    subject_n = norm(subject)
    body_n = norm(body)
    sender_n = norm(sender)
    text_n = norm(f"{subject} {body}")

    # Strong rejection/disposition signals. Keep this BEFORE confirmation and
    # recruiter language. These phrases are intentionally body-aware.
    rejection_phrases = (
        "decided not to move forward",
        "not to move forward with your candidacy",
        "not be moving forward with your candidacy",
        "will not be moving forward with your candidacy",
        "will not move forward with your candidacy",
        "not moving forward with your application",
        "not be moving forward with your application",
        "will not be moving forward with your application",
        "not moving forward with your candidacy",
        "regret to inform you",
        "unfortunately we will not be moving forward",
        "unfortunately we are not moving forward",
        "unfortunately we have decided not to move forward",
        "unfortunately we are unable to move forward",
        "unable to move forward with your candidacy",
        "unable to move forward with your application",
        "we will not be proceeding",
        "we are not proceeding",
        "not selected",
        "you were not selected",
        "not selected for this position",
        "not selected for the role",
        "not selected to move forward",
        "we have decided not to proceed",
        "we have decided to proceed with other candidates",
        "we have decided to move forward with other candidates",
        "we've decided to move forward with other candidates",
        "we’ve decided to move forward with other candidates",
        "decided to move forward with other candidates",
        "move forward with other candidates",
        "moving forward with other candidates",
        "move forward with candidates whose qualifications",
        "moving forward with candidates whose qualifications",
        "pursue other candidates",
        "unable to offer you a position",
        "unable to offer you employment",
        "will not be moving forward",
        "will not move forward",
        "unable to move forward",
        "unable to move forward at this time",
        "are not moving forward",
        "not be progressing",
        "not progressing",
        "other candidates whose qualifications more closely match",
        "other candidates more closely match",
        "not a match at this time",
        "not the right match at this time",
        "we won't be moving forward",
        "we will keep your resume on file",
        "keep your resume on file for future opportunities",
        "wish you the best in your job search",
        "wish you success in your job search",
        "we are unable to offer you the position",
    )
    hit = contains_any(text_n, rejection_phrases)
    if hit:
        return Classification("Rejection", 0.98, f"Body/subject contains rejection signal: '{hit}'.")

    # Defensive fallback for ATS rejection emails that vary contractions and
    # pronouns around the same meaning, e.g. "After thoughtful consideration,
    # we've decided to move forward with other candidates...".
    if ("decided to move forward" in text_n and "other candidate" in text_n):
        return Classification("Rejection", 0.98, "Body contains rejection signal: decided to move forward with other candidates.")
    if ("move forward" in text_n and "other candidate" in text_n and ("after careful consideration" in text_n or "after thoughtful consideration" in text_n)):
        return Classification("Rejection", 0.97, "Body contains rejection signal: moving forward with other candidates after consideration.")
    if ("profile will remain" in text_n and "database" in text_n and "suitable opportunity" in text_n):
        return Classification("Rejection", 0.96, "Body says profile will remain on file/database for future opportunities.")

    if ("unable to move forward" in text_n and ("application" in text_n or "hiring needs" in text_n or "qualifications" in text_n)):
        return Classification("Rejection", 0.98, "Body says employer is unable to move forward with the application.")
    if ("unable to proceed" in text_n and ("application" in text_n or "candidate" in text_n)):
        return Classification("Rejection", 0.97, "Body says employer is unable to proceed with the application.")

    offer_phrases = (
        "pleased to offer you",
        "we are pleased to offer you",
        "would like to offer you the position",
        "formal offer of employment",
        "your offer letter",
        "attached is your offer letter",
    )
    hit = contains_any(text_n, offer_phrases)
    if hit:
        return Classification("Offer", 0.98, f"Contains offer signal: '{hit}'.")

    withdrawal_phrases = (
        "withdraw your application",
        "application has been withdrawn",
        "you have withdrawn",
    )
    hit = contains_any(text_n, withdrawal_phrases)
    if hit:
        candidate_initiated = (
            "you withdrew" in text_n
            or "you have withdrawn" in text_n
            or "your withdrawal" in text_n
            or "withdrawn your application" in text_n
            or "request to withdraw" in text_n
            or "i withdraw" in text_n
            or "i am withdrawing" in text_n
            or "i have withdrawn" in text_n
            or "candidate withdrew" in text_n
        )
        if candidate_initiated:
            return Classification("Withdrawal", 0.95, f"Contains candidate-initiated withdrawal signal: '{hit}'.")
        # Do not let employer phrases like "unable to move forward" or generic
        # "withdrawn from consideration" outrank rejection. In ambiguous cases,
        # withdrawal is too rare and should fall through to other classifiers.
    interview_phrases = (
        "interview invitation",
        "invite you to interview",
        "schedule an interview",
        "would like to interview",
        "meet with the team",
        "meet the team",
        "next round",
        "next step is an interview",
        "phone interview",
        "video interview",
    )
    hit = contains_any(text_n, interview_phrases)
    if hit:
        return Classification("Interview Invite", 0.93, f"Contains interview signal: '{hit}'.")

    assessment_phrases = (
        "assessment",
        "take home",
        "take-home",
        "skills test",
        "case study",
        "assignment",
        "coding challenge",
    )
    hit = contains_any(text_n, assessment_phrases)
    if hit:
        return Classification("Assessment", 0.90, f"Contains assessment signal: '{hit}'.")

    scheduling_phrases = (
        "please share your availability",
        "send your availability",
        "calendar invite",
        "schedule a time",
        "scheduling link",
        "calendly",
        "available times",
    )
    hit = contains_any(text_n, scheduling_phrases)
    if hit:
        return Classification("Scheduling", 0.88, f"Contains scheduling signal: '{hit}'.")

    screening_phrases = (
        "screening call",
        "introductory call",
        "recruiter screen",
        "talent screen",
        "initial screen",
    )
    hit = contains_any(text_n, screening_phrases)
    if hit:
        return Classification("Screening Request", 0.88, f"Contains screening signal: '{hit}'.")

    confirmation_phrases = (
        "thank you for submitting your application",
        "thanks for submitting your application",
        "submitted your application",
        "your application was submitted",
        "your application has been submitted",
        "your application with",
        "we received your application",
        "we have received your application",
        "we've received your application",
        "application received",
        "application has been received",
        "application was received",
        "successfully received your application",
        "thank you for applying",
        "thanks for applying",
        "thank you for your application",
        "we appreciate your application",
        "appreciate your application",
        "we received your resume",
        "we've received your resume",
        "we have received your resume",
        "received your resume",
        "resume received",
        "application will be reviewed",
        "will be reviewing your application",
    )
    hit = contains_any(text_n, confirmation_phrases)
    if hit:
        return Classification("Application Confirmation", 0.90, f"Contains confirmation signal: '{hit}'.")

    # Workday/ATS confirmation emails often have short subjects such as
    # "Your Application with Novanta" while the body is sparse or templated.
    # This fallback still comes AFTER rejection/offer/interview checks.
    if subject_n.startswith("your application with") or subject_n.startswith("thank you for submitting your application"):
        return Classification("Application Confirmation", 0.88, "Subject is an ATS application-submission confirmation.")
    if "application" in subject_n and "thank you for submitting" in text_n:
        return Classification("Application Confirmation", 0.88, "Body thanks candidate for submitting the application.")

    # ATS status/update messages from Workday/Ashby/etc. often contain generic
    # signature text like "Talent Acquisition Team." That signature alone should
    # NOT turn the message into Recruiter Outreach. If no stronger outcome above
    # matched, classify common ATS application-status receipts as confirmation/status evidence.
    ats_sender_markers = (
        "myworkday.com",
        "workday",
        "ashbyhq.com",
        "greenhouse",
        "lever.co",
        "smartrecruiters",
        "icims",
        "successfactors",
        "taleo",
    )
    ats_like_sender = any(m in sender_n for m in ats_sender_markers)
    ats_confirmation_subjects = (
        "your application to",
        "your application with",
        "your recent job application",
        "thank you for your interest in",
        "thanks for your interest in",
        "thank you for applying",
        "thanks for applying",
        "application received",
        "application update",
    )
    if (ats_like_sender or "talent acquisition team" in text_n) and contains_any(subject_n, ats_confirmation_subjects):
        if (
            "your application will be taken into careful consideration" in text_n
            or "application will be taken into careful consideration" in text_n
            or "your application" in subject_n
            or "your recent job application" in subject_n
            or "thank you for your interest" in text_n
        ):
            return Classification("Application Confirmation", 0.86, "ATS application-status message; no stronger rejection/interview/offer outcome found.")

    if "your application will be taken into careful consideration" in text_n or "application will be taken into careful consideration" in text_n:
        return Classification("Application Confirmation", 0.86, "Body says the application will be taken into careful consideration.")

    followup_phrases = (
        "following up on your application",
        "follow up on your application",
        "follow-up on your application",
        "checking in on your application",
        "status of your application",
    )
    hit = contains_any(text_n, followup_phrases)
    if hit:
        return Classification("Follow-up", 0.78, f"Contains follow-up signal: '{hit}'.")

    admin_phrases = (
        "verify your candidate account",
        "verify candidate account",
        "activate your candidate account",
        "candidate account verification",
        "confirm your identity",
        "reset your password",
        "candidate account",
        "verification code",
    )
    hit = contains_any(text_n, admin_phrases)
    if hit:
        return Classification("Administrative", 0.90, f"Contains administrative/candidate-account signal: '{hit}'.")

    # Recruiter Outreach is intentionally late and low-confidence. It should
    # capture true outreach/networking messages, not ATS status emails whose only
    # recruiter signal is a generic signature.
    recruiter_phrases = (
        "recruiter",
        "talent acquisition",
        "hiring team",
        "sourcer",
        "recruiting team",
        "open role",
        "opportunity at",
    )
    hit = contains_any(text_n, recruiter_phrases)
    if hit:
        return Classification("Recruiter Outreach", 0.72, f"Contains recruiter/talent signal without stronger outcome: '{hit}'.")

    return Classification("Other", 0.40, "No strong employment lifecycle signal found in subject/body.")


def is_administrative_account_email(subject: str, body: str) -> bool:
    text_n = norm(f"{subject} {body}")
    if classify(subject, body).kind in ARCHIVABLE_KINDS:
        return False
    admin_subjects = (
        "verify your candidate account",
        "verify candidate account",
        "activate your candidate account",
        "candidate account verification",
        "verify your email",
        "verify your email address",
        "activate your account",
        "set your password",
        "reset your password",
    )
    return bool(contains_any(text_n, admin_subjects))


def is_non_substantive_recruiting_email(subject: str, body: str) -> bool:
    text_n = norm(f"{subject} {body}")
    phrases = (
        "how was your recruiting experience",
        "candidate experience survey",
        "recruiting experience survey",
        "tell us about your recruiting experience",
        "share feedback on your recruiting experience",
        "take our candidate survey",
    )
    return bool(contains_any(text_n, phrases))


def has_strong_job_context(text: str) -> bool:
    text_n = norm(text)
    phrases = (
        "application", "applied", "applying", "candidate", "interview",
        "recruiter", "recruiting", "talent acquisition", "hiring team",
        "hiring manager", "screening call", "assessment", "position at",
        "position with", "role at", "role with", "opportunity at",
    )
    return bool(contains_any(text_n, phrases))


def message_is_on_or_after_application(app: dict, msg_dt: datetime) -> bool:
    applied = app.get("date_applied")
    if not applied:
        return True
    try:
        applied_date = applied.date() if hasattr(applied, "date") else applied
        return msg_dt.date() >= applied_date
    except Exception:
        return True


def company_identity_match(app: dict, subject: str, participants: str) -> bool:
    company = norm(app["company"])
    if not company:
        return False
    participant_n = norm(participants)
    subject_n = norm(subject)
    company_pattern = r"(?<![a-z0-9])" + re.escape(company) + r"(?![a-z0-9])"
    if re.search(company_pattern, participant_n) or re.search(company_pattern, subject_n):
        return True
    company_compact = compact_norm(app["company"])
    participant_compact = compact_norm(participants)
    return len(company_compact) >= 5 and company_compact in participant_compact


def load_applications() -> list[dict]:
    wb = load_workbook(LOG_FILE, read_only=True, data_only=False)
    try:
        ws = wb["Applications"]
        h = ws_headers(ws)
        c_id = require_col(h, "Application ID")
        c_folder = require_col(h, "Application Folder")
        c_company = require_col(h, "Company")
        c_role = require_col(h, "Role")
        c_req = find_col(h, "Requisition ID", "Req ID", "ReqId")
        c_url = find_col(h, "Job URL", "URL")
        c_status = find_col(h, "Status")
        c_date = find_col(h, "Date Applied")
        rows = []
        rownum = 1
        for values in ws.iter_rows(min_row=2, values_only=True):
            rownum += 1
            app_id = values[c_id - 1] if len(values) >= c_id else None
            folder = values[c_folder - 1] if len(values) >= c_folder else None
            company = values[c_company - 1] if len(values) >= c_company else None
            role = values[c_role - 1] if len(values) >= c_role else None
            if not app_id or not folder or not company or not role:
                continue
            rows.append({
                "row": rownum,
                "id": str(app_id).strip(),
                "company": str(company).strip(),
                "role": str(role).strip(),
                "reqid": str(values[c_req - 1] if c_req and len(values) >= c_req and values[c_req - 1] else "").strip(),
                "url": str(values[c_url - 1] if c_url and len(values) >= c_url and values[c_url - 1] else "").strip(),
                "folder": Path(str(folder)),
                "status": str(values[c_status - 1] if c_status and len(values) >= c_status and values[c_status - 1] else "").strip(),
                "date_applied": values[c_date - 1] if c_date and len(values) >= c_date else None,
            })
        return rows
    finally:
        wb.close()


def application_score(app: dict, subject: str, body: str, participants: str, thread_id: str, thread_map: dict, msg_dt: datetime) -> int:
    if thread_map.get(thread_id) == app["id"]:
        return 100
    if not message_is_on_or_after_application(app, msg_dt):
        return 0
    subject_n = norm(subject)
    combined_n = norm(f"{subject} {body}")
    combined_compact = compact_norm(f"{subject} {body}")
    req = norm(app["reqid"])
    req_compact = compact_norm(app["reqid"])
    role = norm(app["role"])
    if req and (req in combined_n or (req_compact and len(req_compact) >= 5 and req_compact in combined_compact)):
        return 95
    employer_identified = company_identity_match(app, subject, participants)
    job_context = has_strong_job_context(f"{subject} {body}")
    if role and role in subject_n and job_context:
        return 92
    if employer_identified and role and role in combined_n and job_context:
        return 90
    if employer_identified and job_context:
        return 80
    return 0


def choose_application(apps: list[dict], subject: str, body: str, participants: str, thread_id: str, thread_map: dict, msg_dt: datetime) -> dict | None:
    scored = sorted(((application_score(a, subject, body, participants, thread_id, thread_map, msg_dt), a) for a in apps), key=lambda x: x[0], reverse=True)
    if scored and scored[0][0] >= 80:
        best_score, best = scored[0]
        if len(scored) == 1 or scored[1][0] != best_score or best_score == 100:
            return best
        return None
    combined_n = norm(f"{subject} {body}")
    if has_strong_job_context(combined_n):
        role_matches = [a for a in apps if message_is_on_or_after_application(a, msg_dt) and norm(a["role"]) and norm(a["role"]) in combined_n]
        if len(role_matches) == 1:
            return role_matches[0]
    return None


def render_pdf(html_body: str, headers: dict[str, str], output_path: Path, gmail_id: str, classification: Classification):
    header_html = f"""
    <div style="font-family:Segoe UI,Arial,sans-serif;font-size:11px;border-bottom:1px solid #aaa;padding-bottom:10px;margin-bottom:18px">
      <div><b>Date:</b> {html_lib.escape(headers.get('date',''))}</div>
      <div><b>From:</b> {html_lib.escape(headers.get('from',''))}</div>
      <div><b>To:</b> {html_lib.escape(headers.get('to',''))}</div>
      <div><b>Subject:</b> {html_lib.escape(headers.get('subject',''))}</div>
      <div><b>Gmail Message ID:</b> {html_lib.escape(gmail_id)}</div>
      <div><b>Classification:</b> {html_lib.escape(classification.kind)} ({classification.confidence:.2f}) — {html_lib.escape(classification.reason)}</div>
    </div>
    """
    doc_html = f"""<!doctype html><html><head><meta charset="utf-8"><style>
        body {{ font-family: Segoe UI, Arial, sans-serif; font-size: 11pt; color:#111; margin: 28px; }}
        img {{ max-width: 100%; }} table {{ max-width: 100%; }} a {{ color:#0645ad; }}
      </style></head><body>{header_html}{html_body}</body></html>"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        browser = None
        last_error = None
        for channel in ("msedge", "chrome"):
            try:
                browser = p.chromium.launch(channel=channel, headless=True)
                break
            except Exception as e:
                last_error = e
        if browser is None:
            try:
                browser = p.chromium.launch(headless=True)
            except Exception:
                raise RuntimeError(f"Could not launch browser through Playwright. Last error: {last_error}")
        page = browser.new_page()
        page.set_content(doc_html, wait_until="load")
        page.emulate_media(media="screen")
        page.pdf(path=str(output_path), format="Letter", print_background=True, margin={"top": "0.45in", "right": "0.45in", "bottom": "0.45in", "left": "0.45in"})
        browser.close()


def unique_pdf_path(folder: Path, dt: datetime, kind: str, direction: str) -> Path:
    stem = f"{dt:%Y%m%d_%H%M}_{safe_filename(kind)}_{direction}"
    path = folder / f"{stem}.pdf"
    n = 2
    while path.exists():
        path = folder / f"{stem}_{n}.pdf"
        n += 1
    return path


def backup_workbook(label: str) -> Path:
    backup_dir = ROOT / "Backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"Job_Search_Log_before_{label}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.copy2(LOG_FILE, backup)
    return backup


def find_existing_correspondence(corr_ws, h: dict[str, int], app_id: str, gmail_id: str, dt: datetime, subject: str, direction: str) -> int | None:
    c_app = require_col(h, "Application ID")
    c_gmail = find_col(h, "Gmail Message ID")
    c_date = find_col(h, "Date", "Correspondence Date")
    c_subject = find_col(h, "Subject")
    c_direction = find_col(h, "Direction")
    subj_n = norm(subject)
    for r in range(2, corr_ws.max_row + 1):
        if str(corr_ws.cell(r, c_app).value or "").strip() != app_id:
            continue
        if c_gmail and gmail_id and str(corr_ws.cell(r, c_gmail).value or "").strip() == gmail_id:
            return r
        if c_subject and c_direction:
            row_subject = norm(corr_ws.cell(r, c_subject).value)
            row_direction = str(corr_ws.cell(r, c_direction).value or "").strip()
            if row_subject == subj_n and row_direction == direction:
                if c_date:
                    row_dt = corr_ws.cell(r, c_date).value
                    try:
                        if hasattr(row_dt, "date") and row_dt.date() == dt.date():
                            return r
                    except Exception:
                        pass
                else:
                    return r
    return None


def existing_pdf_for_message(app: dict, gmail_id: str, dt: datetime, subject: str, direction: str) -> Path | None:
    if not LOG_FILE.exists():
        return None
    wb = load_workbook(LOG_FILE, read_only=True, data_only=False)
    try:
        if "Correspondence" not in wb.sheetnames:
            return None
        ws = wb["Correspondence"]
        h = ws_headers(ws)
        c_pdf = find_col(h, "PDF Filename", "File", "Filename")
        if not c_pdf:
            return None
        row = find_existing_correspondence(ws, h, app["id"], gmail_id, dt, subject, direction)
        if not row:
            return None
        name = str(ws.cell(row, c_pdf).value or "").strip()
        if not name:
            return None
        candidate = app["folder"] / name
        return candidate if candidate.exists() else None
    finally:
        wb.close()


def write_or_update_excel(app: dict, dt: datetime, direction: str, sender_to: str, subject: str, classification: Classification, pdf_path: Path, gmail_id: str, reprocess: bool) -> str:
    wb = load_workbook(LOG_FILE)
    try:
        apps_ws = wb["Applications"]
        corr_ws = wb["Correspondence"]
        app_h = ws_headers(apps_ws)
        corr_h = ws_headers(corr_ws)
        c_gmail = ensure_col(corr_ws, corr_h, "Gmail Message ID")
        c_class_reason = ensure_col(corr_ws, corr_h, "Classification Reason")
        c_class_conf = ensure_col(corr_ws, corr_h, "Classification Confidence")

        c_app = require_col(corr_h, "Application ID")
        c_date = find_col(corr_h, "Date", "Correspondence Date")
        c_company = find_col(corr_h, "Company")
        c_role = find_col(corr_h, "Role")
        c_direction = find_col(corr_h, "Direction")
        c_sender = find_col(corr_h, "Sender/To", "Contact", "Sender", "Recipient")
        c_subject = find_col(corr_h, "Subject")
        c_kind = require_col(corr_h, "Correspondence Type", "Type")
        c_pdf = find_col(corr_h, "PDF Filename", "File", "Filename")
        c_folder = find_col(corr_h, "Application Folder", "Folder")
        c_notes = find_col(corr_h, "Notes")

        row = find_existing_correspondence(corr_ws, corr_h, app["id"], gmail_id, dt, subject, direction)
        action = "updated" if row else "inserted"
        if not row:
            row = first_blank_row(corr_ws, c_app)

        existing_kind = str(corr_ws.cell(row, c_kind).value or "").strip()
        if row and existing_kind and existing_kind != classification.kind and c_pdf:
            old_pdf_name = str(corr_ws.cell(row, c_pdf).value or "").strip()
            old_pdf = app["folder"] / old_pdf_name if old_pdf_name else None
            if old_pdf and old_pdf.exists():
                new_pdf = app["folder"] / old_pdf.name.replace(safe_filename(existing_kind), safe_filename(classification.kind))
                if new_pdf != old_pdf and not new_pdf.exists():
                    old_pdf.rename(new_pdf)
                    pdf_path = new_pdf

        values = {
            c_app: app["id"],
            c_gmail: gmail_id,
            c_class_reason: classification.reason,
            c_class_conf: classification.confidence,
        }
        if c_date: values[c_date] = dt.replace(tzinfo=None)
        if c_company: values[c_company] = app["company"]
        if c_role: values[c_role] = app["role"]
        if c_direction: values[c_direction] = direction
        if c_sender: values[c_sender] = sender_to
        if c_subject: values[c_subject] = subject
        if c_kind: values[c_kind] = classification.kind
        if c_pdf: values[c_pdf] = pdf_path.name
        if c_folder: values[c_folder] = str(app["folder"])
        if c_notes and reprocess:
            note = str(corr_ws.cell(row, c_notes).value or "").strip()
            suffix = f"Reclassified {datetime.now():%Y-%m-%d %H:%M} by content-aware watcher."
            if suffix not in note:
                values[c_notes] = (note + " | " + suffix).strip(" |")
        for col, value in values.items():
            corr_ws.cell(row, col).value = value

        app_row = app["row"]
        c_latest_date = find_col(app_h, "Latest Correspondence Date")
        c_latest_type = find_col(app_h, "Latest Correspondence Type")
        if c_latest_date: apps_ws.cell(app_row, c_latest_date).value = dt.replace(tzinfo=None)
        if c_latest_type: apps_ws.cell(app_row, c_latest_type).value = classification.kind

        c_status = find_col(app_h, "Status")
        c_conf = find_col(app_h, "Confirmation Received")
        c_interview = find_col(app_h, "Interview Date")
        c_rej = find_col(app_h, "Rejection Date")
        if direction == "Incoming" and c_status:
            if classification.kind == "Application Confirmation":
                if c_conf: apps_ws.cell(app_row, c_conf).value = "Yes"
                current = str(apps_ws.cell(app_row, c_status).value or "")
                if current in ("Qualified", "Prepared", "Resume Created", "Application In Progress", "Ready for Review"):
                    apps_ws.cell(app_row, c_status).value = "Applied"
            elif classification.kind in ("Recruiter Outreach", "Screening Request"):
                current = str(apps_ws.cell(app_row, c_status).value or "")
                if current not in ("Applied", "Interview", "Interviewing", "Rejected", "Offer", "Withdrawn", "Closed", "Skipped", "Do Not Reapply"):
                    apps_ws.cell(app_row, c_status).value = "Recruiter Contact"
            elif classification.kind in ("Interview Invite", "Scheduling", "Assessment"):
                apps_ws.cell(app_row, c_status).value = "Interviewing"
                if classification.kind == "Interview Invite" and c_interview and not apps_ws.cell(app_row, c_interview).value:
                    apps_ws.cell(app_row, c_interview).value = dt.replace(tzinfo=None)
            elif classification.kind == "Rejection":
                apps_ws.cell(app_row, c_status).value = "Rejected"
                if c_rej: apps_ws.cell(app_row, c_rej).value = dt.date()
            elif classification.kind == "Offer":
                apps_ws.cell(app_row, c_status).value = "Offer"
            elif classification.kind == "Withdrawal":
                apps_ws.cell(app_row, c_status).value = "Withdrawn"

        tmp = LOG_FILE.with_name(f"{LOG_FILE.stem}.gmail_archive.{datetime.now():%Y%m%d%H%M%S%f}.tmp.xlsx")
        wb.save(tmp)
        wb.close()
        check = load_workbook(tmp, read_only=True, data_only=False)
        check.close()
        tmp.replace(LOG_FILE)
        return action
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        raise


def iter_gmail_messages(service, query: str, *, page_size: int = 25, limit: int = 0, max_retries: int = 5, base_sleep: float = 8.0):
    page_token = None
    seen = 0
    page_size = max(1, min(int(page_size or 25), 100))
    while True:
        req = service.users().messages().list(userId="me", q=query, maxResults=page_size, pageToken=page_token)
        result = execute_gmail_request(req, label="messages.list", max_retries=max_retries, base_sleep=base_sleep)
        for item in result.get("messages", []):
            yield item
            seen += 1
            if limit and seen >= limit:
                return
        page_token = result.get("nextPageToken")
        if not page_token:
            break


def process_messages(args) -> None:
    apps = load_applications()
    if not apps:
        print("No application rows with folders found. Nothing to archive.")
        return

    service = get_service()
    processed = set(load_json(PROCESSED_FILE, []))
    thread_map = load_json(THREAD_MAP_FILE, {})
    pending = load_json(PENDING_FILE, [])

    if args.message_id:
        message_refs = [{"id": args.message_id}]
        query = f"id:{args.message_id}"
    else:
        query = args.query or f"in:anywhere newer_than:{args.days}d -category:promotions"
        message_refs = list(iter_gmail_messages(
            service,
            query,
            page_size=args.page_size,
            limit=args.limit,
            max_retries=args.max_retries,
            base_sleep=args.backoff_seconds,
        ))

    print(f"Gmail query: {query}")
    print(f"Candidate messages: {len(message_refs)}")
    if args.reprocess:
        print("Mode: REPROCESS processed messages too; duplicate-safe update/insert enabled.")
    if args.dry_run:
        print("Mode: DRY RUN; no PDFs/workbook/state changes.")
    if args.limit:
        print(f"Mode: LIMIT {args.limit} message(s).")
    if args.sleep_seconds:
        print(f"Mode: throttle {args.sleep_seconds:.2f}s after each full-message read.")
    if args.direction != "all":
        print(f"Mode: direction filter {args.direction}.")

    if not args.dry_run:
        backup = backup_workbook("gmail_content_aware")
        print(f"Workbook backup: {backup}")

    reviewed = archived = updated = skipped_processed = skipped_other = unmatched = 0

    for item in message_refs:
        msg_id = item["id"]
        if msg_id in processed and not args.reprocess:
            skipped_processed += 1
            continue
        try:
            req = service.users().messages().get(userId="me", id=msg_id, format="full")
            msg = execute_gmail_request(req, label=f"messages.get {msg_id}", max_retries=args.max_retries, base_sleep=args.backoff_seconds)
            if args.sleep_seconds:
                time.sleep(args.sleep_seconds)
            payload = msg.get("payload", {})
            headers = get_headers(payload)
            msg_direction = direction_from_headers(headers)
            if args.direction != "all" and msg_direction.lower() != args.direction:
                skipped_other += 1
                print(f"SKIP DIRECTION {msg_id} | {msg_direction}")
                continue
            html_body, text_body, visible_text = message_content(payload)
            subject = headers.get("subject", "")
            sender = headers.get("from", "")
            to = headers.get("to", "")
            thread_id = msg.get("threadId", "")
            dt = email_datetime(headers)
            participants = f"{sender} {to}"
            app = choose_application(apps, subject, visible_text, participants, thread_id, thread_map, dt)
            if not app:
                unmatched += 1
                continue
            classification = classify(subject, visible_text, sender)
            if is_administrative_account_email(subject, visible_text):
                print(f"SKIP ADMIN {app['id']} | {app['company']} | {subject}")
                if not args.dry_run:
                    processed.add(msg_id)
                    save_json(PROCESSED_FILE, sorted(processed))
                continue
            if is_non_substantive_recruiting_email(subject, visible_text):
                print(f"SKIP SURVEY {app['id']} | {app['company']} | {subject}")
                if not args.dry_run:
                    processed.add(msg_id)
                    save_json(PROCESSED_FILE, sorted(processed))
                continue
            direction = "Outgoing" if "SENT" in (msg.get("labelIds") or []) else "Incoming"
            if classification.kind == "Other":
                print(f"SKIP OTHER {app['id']} | {app['company']} | {direction} | {subject} | {classification.reason}")
                skipped_other += 1
                if not args.dry_run and args.mark_other_processed:
                    processed.add(msg_id)
                    save_json(PROCESSED_FILE, sorted(processed))
                continue
            reviewed += 1
            print(f"MATCH {app['id']} | {app['company']} | {classification.kind} | {direction} | {subject} | {classification.reason}")
            if args.dry_run:
                continue
            app["folder"].mkdir(parents=True, exist_ok=True)
            pdf_path = existing_pdf_for_message(app, msg_id, dt, subject, direction)
            if pdf_path is None:
                pdf_path = unique_pdf_path(app["folder"], dt, classification.kind, direction)
                render_pdf(html_body, headers, pdf_path, msg_id, classification)
            sender_to = to if direction == "Outgoing" else sender
            action = write_or_update_excel(app, dt, direction, sender_to, subject, classification, pdf_path, msg_id, args.reprocess)
            if action == "updated":
                updated += 1
            else:
                archived += 1
            thread_map[thread_id] = app["id"]
            processed.add(msg_id)
            save_json(THREAD_MAP_FILE, thread_map)
            save_json(PROCESSED_FILE, sorted(processed))
        except Exception as exc:
            print(f"ERROR {msg_id}: {exc}")
            if not args.dry_run:
                pending.append({"gmail_id": msg_id, "error": str(exc), "ts": datetime.now().isoformat()})
                save_json(PENDING_FILE, pending)

    print(f"Matched {reviewed} message(s); inserted {archived}; updated {updated}; skipped processed {skipped_processed}; skipped other {skipped_other}; unmatched {unmatched}.")


def main(argv=None):
    parser = argparse.ArgumentParser(description="Archive job-related Gmail (read-only) into application folders.")
    parser.add_argument("--days", type=int, default=14, help="Gmail lookback window; default 14 days.")
    parser.add_argument("--query", default="", help="Override Gmail query. Example: 'in:anywhere after:2026/08/01 -category:promotions'")
    parser.add_argument("--message-id", default="", help="Process one Gmail message ID.")
    parser.add_argument("--dry-run", action="store_true", help="Match/classify but do not save PDFs or update Excel/state.")
    parser.add_argument("--reprocess", action="store_true", help="Process messages even if already listed in gmail_processed.json; duplicate-safe update/insert.")
    parser.add_argument("--mark-other-processed", action="store_true", help="In apply mode, mark matched Other emails processed. Off by default.")
    parser.add_argument("--limit", type=int, default=0, help="Maximum number of candidate messages to fetch/process. Default 0 = no explicit limit.")
    parser.add_argument("--page-size", type=int, default=25, help="Gmail list page size. Smaller values reduce burst quota risk. Default 25; max 100.")
    parser.add_argument("--sleep-seconds", type=float, default=1.0, help="Seconds to sleep after each full-message read. Default 1.0.")
    parser.add_argument("--max-retries", type=int, default=6, help="Retry count for Gmail 403/429 rate-limit errors. Default 6.")
    parser.add_argument("--backoff-seconds", type=float, default=10.0, help="Initial backoff seconds for Gmail rate-limit retries. Default 10.")
    parser.add_argument("--direction", choices=["all", "incoming", "outgoing"], default="all", help="Only process incoming, outgoing, or all messages. Default all.")
    args = parser.parse_args(argv)
    process_messages(args)


if __name__ == "__main__":
    main()
