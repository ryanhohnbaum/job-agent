"""Application tracker commands: check-duplicate, create, mark-applied, set-status, queue, log-search.

All workbook writes: refuse if the file is open in Excel, back up first, write by header name,
save to a temp file, verify, then atomically replace.
"""
from __future__ import annotations

import argparse
import json
import shutil
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

from common import (
    clean_name, first_blank_row, headers, iter_value_rows, jobagent_root, normalized,
    require_workbook_writable, row_value, save_workbook_verified, verify_application_values,
    workbook_path,
)
from schema import ACTIONABLE, APPLICATION_HEADERS, SEARCH_ACTIVITY_HEADERS, STATUSES


# ---------------------------------------------------------------- duplicates

def scan_duplicates(ws, h, company: str, role: str, reqid: str):
    """Return (exact, similar). Any existing status counts: a role already seen is not new."""
    company_n, role_n, req_n = normalized(company), normalized(role), normalized(reqid)
    exact, similar = [], []
    for row, values in iter_value_rows(ws):
        app_id = str(row_value(values, h["Application ID"]) or "").strip()
        if not app_id:
            continue
        r_company = normalized(row_value(values, h["Company"]))
        r_role = normalized(row_value(values, h["Role"]))
        r_req = normalized(row_value(values, h["Requisition ID"]))
        info = {
            "application_id": app_id, "row": row, "status": row_value(values, h["Status"]),
            "company": row_value(values, h["Company"]), "role": row_value(values, h["Role"]),
        }
        if r_company != company_n:
            continue
        if req_n and r_req == req_n:
            exact.append({**info, "basis": "company + requisition"})
        elif not req_n and r_role == role_n:
            exact.append({**info, "basis": "company + role"})
        elif SequenceMatcher(None, r_role, role_n).ratio() >= 0.85:
            similar.append({**info, "basis": "same company, similar title"})
    return exact, similar


def cmd_check_duplicate(args) -> None:
    wb = load_workbook(workbook_path(), read_only=True)
    try:
        ws = wb["Applications"]
        h = headers(ws, APPLICATION_HEADERS)
        exact, similar = scan_duplicates(ws, h, args.company, args.role, args.reqid)
    finally:
        wb.close()
    if args.json:
        print(json.dumps({"exact": exact, "similar": similar}, indent=2, default=str))
    elif exact:
        for m in exact:
            print(f"DUPLICATE | {m['application_id']} | {m['status']} | {m['basis']}")
    elif similar:
        for m in similar:
            print(f"REVIEW | {m['application_id']} | {m['status']} | {m['role']} | {m['basis']}")
    else:
        print("NEW | no matching application found")


# ---------------------------------------------------------------- create

def next_id(ws, id_column: int, ymd: str) -> str:
    prefix = f"APP-{ymd}-"
    numbers = []
    for _, values in iter_value_rows(ws):
        value = row_value(values, id_column)
        if isinstance(value, str) and value.startswith(prefix):
            try:
                numbers.append(int(value.rsplit("-", 1)[1]))
            except ValueError:
                continue
    return f"{prefix}{(max(numbers) + 1 if numbers else 1):03d}"


def write_apply_url(folder: Path, url: str) -> None:
    if not url:
        return
    (folder / "Apply.url").write_text(f"[InternetShortcut]\nURL={url}\n", encoding="utf-8")


def cmd_create(args) -> None:
    found = datetime.strptime(args.date or datetime.now().strftime("%Y%m%d"), "%Y%m%d")
    ymd = found.strftime("%Y%m%d")
    if args.fit is not None and not 0 <= args.fit <= 10:
        raise ValueError("--fit must be between 0 and 10.")
    resume = Path(args.resume).expanduser() if args.resume else None
    if resume and not resume.exists():
        raise FileNotFoundError(resume)

    book = require_workbook_writable()
    wb = load_workbook(book)
    try:
        ws = wb["Applications"]
        h = headers(ws, APPLICATION_HEADERS)
        exact, similar = scan_duplicates(ws, h, args.company, args.role, args.reqid)
        if exact:
            m = exact[0]
            raise RuntimeError(f"Duplicate: {m['application_id']} ({m['status']}, {m['basis']}). Not created.")
        if similar and not args.allow_similar:
            m = similar[0]
            raise RuntimeError(
                f"Possible duplicate: {m['application_id']} '{m['role']}' ({m['status']}). "
                "Review it, then rerun with --allow-similar if this is a different role."
            )

        app_id = next_id(ws, h["Application ID"], ymd)
        folder_name = f"{ymd}_{clean_name(args.company)}_{clean_name(args.role)}"
        if args.reqid:
            folder_name += f"_{clean_name(args.reqid)}"
        app_folder = jobagent_root() / "Applications" / ymd[:4] / folder_name
        app_folder.mkdir(parents=True, exist_ok=True)

        resume_name = ""
        if resume:
            resume_name = resume.name
            if resume.resolve().parent != app_folder.resolve():
                shutil.copy2(resume, app_folder / resume_name)

        row = first_blank_row(ws, h["Application ID"])
        values = {
            "Application ID": app_id, "Date Found": found, "Date Applied": None,
            "Company": args.company, "Role": args.role, "Requisition ID": args.reqid,
            "Location": args.location, "Work Model": args.work_model, "Job URL": args.url,
            "Salary Range": args.salary, "Fit Score": args.fit, "Resume Filename": resume_name,
            "Application Folder": str(app_folder), "Status": "Qualified",
            "Confirmation Received": "Pending", "Latest Correspondence Date": None,
            "Latest Correspondence Type": "", "Interview Date": None, "Rejection Date": None,
            "Evidence Complete": "No", "Notes": args.notes,
        }
        for name, value in values.items():
            ws.cell(row, h[name]).value = value
        write_apply_url(app_folder, args.url)
        expected = {
            "Company": args.company, "Role": args.role, "Requisition ID": args.reqid,
            "Application Folder": str(app_folder), "Status": "Qualified",
        }
        _, backup = save_workbook_verified(
            wb, "create", lambda p: verify_application_values(p, app_id, expected)
        )
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        raise

    result = {"application_id": app_id, "excel_row": row, "application_folder": str(app_folder), "backup": str(backup)}
    if args.json:
        print(json.dumps(result))
    else:
        print(f"CREATED | {app_id} | row {row}\nFolder: {app_folder}\nBackup: {backup}")


# ---------------------------------------------------------------- status changes

def _update_row(application_id: str, action: str, mutate) -> dict:
    book = require_workbook_writable()
    wb = load_workbook(book)
    try:
        ws = wb["Applications"]
        h = headers(ws, APPLICATION_HEADERS)
        rows = [
            r for r, v in iter_value_rows(ws)
            if str(row_value(v, h["Application ID"]) or "").strip() == application_id
        ]
        if len(rows) != 1:
            raise RuntimeError(f"Expected one row for {application_id}; found {len(rows)}.")
        row = rows[0]
        previous = str(ws.cell(row, h["Status"]).value or "").strip()
        expected = mutate(ws, h, row)
        if expected is None:  # nothing to change
            wb.close()
            return {"result": "unchanged", "application_id": application_id, "status": previous}
        _, backup = save_workbook_verified(wb, action, lambda p: verify_application_values(p, application_id, expected))
        return {"result": action, "application_id": application_id, "previous_status": previous, "backup": str(backup)}
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        raise


def _append_note(cell, text: str) -> None:
    existing = str(cell.value or "").strip()
    if text not in existing:
        cell.value = f"{existing} {text}".strip()


def cmd_mark_applied(args) -> None:
    if not args.confirmed_submitted:
        raise SystemExit("Refusing: pass --confirmed-submitted only after the user confirms the employer form was submitted.")
    applied = datetime.strptime(args.date, "%Y-%m-%d") if args.date else datetime.now()

    def mutate(ws, h, row):
        status = str(ws.cell(row, h["Status"]).value or "").strip()
        if status.casefold() == "applied" and ws.cell(row, h["Date Applied"]).value:
            return None
        ws.cell(row, h["Date Applied"]).value = applied
        ws.cell(row, h["Status"]).value = "Applied"
        if ws.cell(row, h["Confirmation Received"]).value in (None, ""):
            ws.cell(row, h["Confirmation Received"]).value = "Pending"
        _append_note(ws.cell(row, h["Notes"]), f"Marked Applied on {applied:%Y-%m-%d}.")
        return {"Date Applied": applied, "Status": "Applied"}

    result = _update_row(args.application_id.strip(), "mark_applied", mutate)
    print(json.dumps(result) if args.json else _fmt(result))


def cmd_set_status(args) -> None:
    valid = {s.casefold(): s for s in STATUSES}
    if args.status.casefold() not in valid:
        raise SystemExit(f"Unknown status {args.status!r}. Use one of: {', '.join(STATUSES)}")
    if args.status.casefold() == "applied":
        raise SystemExit("Use mark-applied for Applied so the submission confirmation gate is enforced.")
    new_status = valid[args.status.casefold()]

    def mutate(ws, h, row):
        if str(ws.cell(row, h["Status"]).value or "").strip() == new_status and not args.note:
            return None
        ws.cell(row, h["Status"]).value = new_status
        if new_status == "Rejected" and not ws.cell(row, h["Rejection Date"]).value:
            ws.cell(row, h["Rejection Date"]).value = datetime.now()
        if args.note:
            _append_note(ws.cell(row, h["Notes"]), args.note)
        return {"Status": new_status}

    result = _update_row(args.application_id.strip(), "set_status", mutate)
    print(json.dumps(result) if args.json else _fmt(result))


def _fmt(result: dict) -> str:
    return " | ".join(f"{k}={v}" for k, v in result.items())


# ---------------------------------------------------------------- read views

def serializable(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def cmd_queue(args) -> None:
    wb = load_workbook(workbook_path(), read_only=True)
    rows = []
    try:
        ws = wb["Applications"]
        h = headers(ws, APPLICATION_HEADERS)
        for _, values in iter_value_rows(ws):
            app_id = str(row_value(values, h["Application ID"]) or "").strip()
            if not app_id:
                continue
            status = str(row_value(values, h["Status"]) or "").strip()
            if not args.all and status.casefold() not in ACTIONABLE:
                continue
            rows.append({
                "application_id": app_id, "status": status,
                "company": row_value(values, h["Company"]), "role": row_value(values, h["Role"]),
                "requisition_id": row_value(values, h["Requisition ID"]),
                "fit_score": row_value(values, h["Fit Score"]),
                "location": row_value(values, h["Location"]), "work_model": row_value(values, h["Work Model"]),
                "job_url": row_value(values, h["Job URL"]),
                "application_folder": row_value(values, h["Application Folder"]),
                "date_found": serializable(row_value(values, h["Date Found"])),
            })
    finally:
        wb.close()
    rows.sort(key=lambda r: (-(float(r["fit_score"]) if r["fit_score"] not in (None, "") else -1),
                             str(r["company"] or "").casefold()))
    if args.json:
        print(json.dumps(rows, indent=2, ensure_ascii=False))
        return
    if not rows:
        print("Nothing in the queue.")
        return
    for r in rows:
        fit = r["fit_score"] if r["fit_score"] not in (None, "") else "-"
        print(f"{r['application_id']} | {r['status']} | fit {fit} | {r['company']} | {r['role']} | req {r['requisition_id'] or '-'}")
        if r["job_url"]:
            print(f"  Apply:  {r['job_url']}")
        if r["application_folder"]:
            print(f"  Folder: {r['application_folder']}")


def cmd_log_search(args) -> None:
    book = require_workbook_writable()
    wb = load_workbook(book)
    try:
        ws = wb["Search Activity"]
        h = headers(ws, SEARCH_ACTIVITY_HEADERS)
        row = first_blank_row(ws, h["Date"])
        for name, value in {
            "Date": datetime.now(), "Search / Source": args.source, "Roles Reviewed": args.reviewed,
            "Qualified Roles": args.qualified, "Applications Started": args.started,
            "Applications Submitted": args.submitted, "Notes": args.notes,
        }.items():
            ws.cell(row, h[name]).value = value

        def verify(path):
            chk = load_workbook(path, read_only=True)
            try:
                if chk["Search Activity"].cell(row, h["Search / Source"]).value != args.source:
                    raise RuntimeError("Search Activity verification failed.")
            finally:
                chk.close()

        _, backup = save_workbook_verified(wb, "log_search", verify)
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        raise
    print(f"LOGGED | row {row} | backup {backup}")


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="JobAgent application tracker.")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("check-duplicate", help="Read-only duplicate screen against the tracker.")
    p.add_argument("--company", required=True)
    p.add_argument("--role", required=True)
    p.add_argument("--reqid", default="")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_check_duplicate)

    p = sub.add_parser("create", help="Create one application row and folder.")
    p.add_argument("--company", required=True)
    p.add_argument("--role", required=True)
    p.add_argument("--reqid", default="")
    p.add_argument("--url", default="")
    p.add_argument("--location", default="")
    p.add_argument("--work-model", default="")
    p.add_argument("--salary", default="")
    p.add_argument("--fit", type=float)
    p.add_argument("--resume", help="Path to the tailored resume; it is copied into the application folder.")
    p.add_argument("--notes", default="")
    p.add_argument("--date", help="YYYYMMDD; defaults to today")
    p.add_argument("--allow-similar", action="store_true", help="Proceed despite a similar-title match at the same company.")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_create)

    p = sub.add_parser("mark-applied", help="Mark Applied. Only after the user confirms submission.")
    p.add_argument("--application-id", required=True)
    p.add_argument("--confirmed-submitted", action="store_true")
    p.add_argument("--date", help="YYYY-MM-DD; defaults to today")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_mark_applied)

    p = sub.add_parser("set-status", help="Set a non-Applied status (Skipped, On Hold, Interview, ...).")
    p.add_argument("--application-id", required=True)
    p.add_argument("--status", required=True)
    p.add_argument("--note", default="")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_set_status)

    p = sub.add_parser("queue", help="Show actionable applications (Qualified, On Hold).")
    p.add_argument("--all", action="store_true")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_queue)

    p = sub.add_parser("log-search", help="Append a row to the Search Activity sheet.")
    p.add_argument("--source", required=True)
    p.add_argument("--reviewed", type=int, default=0)
    p.add_argument("--qualified", type=int, default=0)
    p.add_argument("--started", type=int, default=0)
    p.add_argument("--submitted", type=int, default=0)
    p.add_argument("--notes", default="")
    p.set_defaults(func=cmd_log_search)

    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
