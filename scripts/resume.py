"""Resume tools: build, dump, apply, check.

  build  master.json -> master.docx     Build a clean, ATS-safe, Calibri master from structured data.
  dump   in.docx                        List body paragraphs with index, style, and formatting hints.
  apply  master.docx edits.json out     Clone the master and apply text edits, keeping formatting.
  check  file.docx                      Render to PDF, count pages, list fonts, print extracted text.

The master resume is the factual source. `apply` edits a clone of it; it never edits in place.
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Calibri"


# ------------------------------------------------------------------ build

def _set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def _para(doc, text="", size=10.5, bold=False, italic=False, align=None, before=0, after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0
    if align:
        p.alignment = align
    if text:
        _set_font(p.add_run(text), size, bold, italic)
    return p


def _bottom_border(p):
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "444444")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    ppr.append(borders)


def _bullet(doc, text, size, text_width):
    p = _para(doc, size=size)
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Inches(0.25), Inches(-0.18)
    pf.tab_stops.add_tab_stop(Inches(0.25))
    _set_font(p.add_run("\u2022\t" + text), size)
    return p


def build(spec: dict, out: Path) -> None:
    body = float(spec.get("body_pt", 10.5))
    margin = float(spec.get("margin_in", 0.7))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(margin)
    sec.top_margin = sec.bottom_margin = Inches(margin)
    width = 8.5 - 2 * margin
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(body)

    center = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, spec["name"], 20, True, align=center, after=1)
    if spec.get("headline"):
        _para(doc, spec["headline"], body + 1, True, align=center, after=1)
    for line in [spec.get("contact", ""), *spec.get("lines", [])]:
        if line:
            _para(doc, line, body - 0.5, align=center)
    if spec.get("skills"):
        _para(doc, spec["skills"], body - 0.5, align=center, before=3)

    for section in spec.get("sections", []):
        head = _para(doc, section["heading"].upper(), body + 0.5, True, before=8, after=3)
        _bottom_border(head)
        for text in section.get("paragraphs", []):
            _para(doc, text, body, after=2)
        for text in section.get("bullets", []):
            _bullet(doc, text, body, width)
        for job in section.get("jobs", []):
            line = _para(doc, size=body, before=4)
            line.paragraph_format.tab_stops.add_tab_stop(Inches(width), WD_TAB_ALIGNMENT.RIGHT)
            employer = job["employer"] + (f", {job['location']}" if job.get("location") else "")
            _set_font(line.add_run(employer), body, True)
            if job.get("dates"):
                _set_font(line.add_run("\t" + job["dates"]), body, True)
            if job.get("title"):
                _para(doc, job["title"], body, italic=True, after=1)
            for text in job.get("bullets", []):
                _bullet(doc, text, body, width)
    doc.save(out)


# ------------------------------------------------------------------ dump

def body_paragraphs(doc):
    """Body-level paragraphs in document order (the indexing used by dump/apply)."""
    return list(doc.paragraphs)


def cmd_dump(args) -> None:
    doc = Document(args.docx)
    rows = []
    for i, p in enumerate(body_paragraphs(doc)):
        runs = [r for r in p.runs if r.text]
        rows.append({
            "index": i, "style": p.style.name, "runs": len(runs),
            "bold_lead": bool(runs and runs[0].bold and len(runs) > 1 and not all(r.bold for r in runs)),
            "text": p.text,
        })
    if args.json:
        print(json.dumps(rows, ensure_ascii=False, indent=2))
        return
    for r in rows:
        flag = " [bold-lead: edit with segments]" if r["bold_lead"] else ""
        print(f"{r['index']:>3} | {r['style']} | {r['text']!r}{flag}")
    if doc.tables:
        print(f"\nNOTE: {len(doc.tables)} table(s) present; their text is not indexed. Tables are discouraged for ATS.")


# ------------------------------------------------------------------ apply

def _write_paragraph(p, value) -> None:
    """value: str (formatting of the first text run) or [[text, bold], ...] segments."""
    src_runs = [r for r in p.runs if r.text] or list(p.runs)
    base = copy.deepcopy(src_runs[0]._element.rPr) if src_runs and src_runs[0]._element.rPr is not None else None
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for hl in p._p.findall(qn("w:hyperlink")):
        p._p.remove(hl)
    segments = [[value, None]] if isinstance(value, str) else value
    for text, bold in segments:
        run = p.add_run(text)
        if base is not None:
            run._element.insert(0, copy.deepcopy(base))
        if bold is not None:
            run.bold = bool(bold)


def apply_edits(master: Path, edits: dict, out: Path) -> None:
    doc = Document(master)
    paras = body_paragraphs(doc)
    n = len(paras)

    def check(idx):
        idx = int(idx)
        if not 0 <= idx < n:
            raise IndexError(f"Paragraph index {idx} out of range (0..{n - 1}). Run `dump` on the master.")
        return idx

    for key, value in (edits.get("replace") or {}).items():
        _write_paragraph(paras[check(key)], value)
    moves = edits.get("move") or []
    for m in moves:
        src, before = paras[check(m["idx"])]._p, paras[check(m["before"])]._p
        before.addprevious(src)
    for idx in edits.get("delete") or []:
        el = paras[check(idx)]._p
        el.getparent().remove(el)
    doc.save(out)


def cmd_apply(args) -> None:
    edits = json.loads(Path(args.edits).read_text(encoding="utf-8"))
    out = Path(args.out)
    if out.resolve() == Path(args.master).resolve():
        raise SystemExit("Refusing to overwrite the master. Write the tailored resume to a new file.")
    if "resume" in out.stem.lower():
        print("WARNING: filename contains 'resume'; the convention is YYYYMMDD_Company_Role.docx", file=sys.stderr)
    out.parent.mkdir(parents=True, exist_ok=True)
    apply_edits(Path(args.master), edits, out)
    print(f"WROTE {out}")


# ------------------------------------------------------------------ check

def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
    ):
        if os.path.exists(candidate):
            return candidate
    return None


def render_pdf(docx: Path, outdir: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice not found. Install it (free) to verify page count: https://www.libreoffice.org/download"
        )
    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx)],
        check=True, capture_output=True, timeout=180,
    )
    pdf = outdir / (docx.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("PDF conversion produced no output.")
    return pdf


def fonts_used(docx: Path) -> set[str]:
    doc = Document(docx)
    names = set()
    default = doc.styles["Normal"].font.name
    for p in doc.paragraphs:
        for r in p.runs:
            if r.text.strip():
                names.add(r.font.name or p.style.font.name or default or "(inherited)")
    return names


def cmd_check(args) -> None:
    from pypdf import PdfReader

    docx = Path(args.docx)
    outdir = Path(args.outdir) if args.outdir else Path(tempfile.mkdtemp(prefix="jobagent_render_"))
    pdf = render_pdf(docx, outdir)
    reader = PdfReader(str(pdf))
    pages = len(reader.pages)
    fonts = fonts_used(docx)
    problems = []
    if args.pages and pages != args.pages:
        problems.append(f"page count is {pages}, expected {args.pages}")
    if fonts - {FONT}:
        problems.append(f"non-{FONT} fonts: {sorted(fonts - {FONT})}")
    if re.search(r"resume", docx.stem, re.I):
        problems.append("filename contains 'resume'")
    print(f"PDF: {pdf}\nPages: {pages}\nFonts: {sorted(fonts)}")
    last = reader.pages[-1].extract_text() or ""
    print(f"Last page text lines: {len([l for l in last.splitlines() if l.strip()])}")
    if args.text:
        for i, page in enumerate(reader.pages, 1):
            print(f"\n--- page {i} (extracted) ---\n{page.extract_text()}")
    print("RESULT: " + ("OK" if not problems else "PROBLEMS: " + "; ".join(problems)))
    print("NOTE: LibreOffice substitutes Carlito for Calibri if Calibri is missing; metrics match, so page count is reliable.")
    sys.exit(1 if problems else 0)


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="JobAgent resume tools.")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("build")
    p.add_argument("spec_json")
    p.add_argument("out_docx")
    p.set_defaults(func=lambda a: (build(json.loads(Path(a.spec_json).read_text(encoding="utf-8")), Path(a.out_docx)),
                                   print(f"WROTE {a.out_docx}")))

    p = sub.add_parser("dump")
    p.add_argument("docx")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_dump)

    p = sub.add_parser("apply")
    p.add_argument("master")
    p.add_argument("edits")
    p.add_argument("out")
    p.set_defaults(func=cmd_apply)

    p = sub.add_parser("check")
    p.add_argument("docx")
    p.add_argument("--pages", type=int, default=0, help="Expected page count (0 = just report). Use the master resume page count.")
    p.add_argument("--outdir", help="Where to write the rendered PDF (default: temp dir).")
    p.add_argument("--text", action="store_true", help="Print extracted plain text (ATS check).")
    p.set_defaults(func=cmd_check)

    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
